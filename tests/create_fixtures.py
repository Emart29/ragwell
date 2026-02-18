"""Create test fixture files for parser tests."""
import fitz  # PyMuPDF
from docx import Document
from docx.shared import Pt
from pathlib import Path
import csv


def create_pdf_fixture(output_path: Path):
    """Create a simple PDF test file with PyMuPDF."""
    doc = fitz.open()
    
    # Page 1: Title and content
    page1 = doc.new_page()
    page1.insert_text((72, 72), "Test Document Title", fontsize=20)
    page1.insert_text((72, 120), "Author: Test Author", fontsize=12)
    page1.insert_text((72, 160), "This is a test PDF document for Ragwell.", fontsize=12)
    page1.insert_text((72, 200), "It contains multiple pages and sample content.", fontsize=12)
    
    # Page 2: Table
    page2 = doc.new_page()
    page2.insert_text((72, 72), "Sample Data Table", fontsize=16)
    
    # Create a simple table using text
    headers = ["Name", "Age", "City"]
    rows = [
        ["Alice", "30", "New York"],
        ["Bob", "25", "San Francisco"],
        ["Charlie", "35", "Chicago"]
    ]
    
    y_pos = 120
    # Headers
    x_pos = 72
    for header in headers:
        page2.insert_text((x_pos, y_pos), header, fontsize=12)
        x_pos += 100
    
    # Rows
    y_pos += 30
    for row in rows:
        x_pos = 72
        for cell in row:
            page2.insert_text((x_pos, y_pos), cell, fontsize=12)
            x_pos += 100
        y_pos += 30
    
    # Add metadata
    doc.set_metadata({
        "title": "Test PDF Document",
        "author": "Test Author",
        "subject": "Ragwell Testing"
    })
    
    doc.save(output_path)
    doc.close()


def create_docx_fixture(output_path: Path):
    """Create a DOCX test file with python-docx."""
    doc = Document()
    
    # Title
    title = doc.add_heading("Test Document Title", level=0)
    
    # Metadata paragraph
    doc.add_paragraph("Author: Test Author")
    doc.add_paragraph("")
    
    # Content
    doc.add_heading("Introduction", level=1)
    doc.add_paragraph("This is a test DOCX document for Ragwell.")
    doc.add_paragraph("It contains structured content with headings and tables.")
    
    doc.add_heading("Section 1: Sample Content", level=1)
    doc.add_paragraph("Here is some sample content for testing.")
    doc.add_paragraph("This content spans multiple paragraphs.")
    
    # Table
    doc.add_heading("Data Table", level=2)
    table = doc.add_table(rows=4, cols=3)
    table.style = 'Table Grid'
    
    # Header row
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Name"
    hdr_cells[1].text = "Age"
    hdr_cells[2].text = "City"
    
    # Data rows
    data = [
        ["Alice", "30", "New York"],
        ["Bob", "25", "San Francisco"],
        ["Charlie", "35", "Chicago"]
    ]
    
    for i, row_data in enumerate(data, 1):
        row_cells = table.rows[i].cells
        for j, value in enumerate(row_data):
            row_cells[j].text = value
    
    # Add core properties
    doc.core_properties.title = "Test DOCX Document"
    doc.core_properties.author = "Test Author"
    doc.core_properties.subject = "Ragwell Testing"
    
    doc.save(output_path)


def create_html_fixture(output_path: Path):
    """Create an HTML test file."""
    html_content = """<!DOCTYPE html>
<html>
<head>
    <title>Test HTML Document</title>
    <meta name="description" content="Test document for Ragwell HTML parser">
    <meta name="author" content="Test Author">
</head>
<body>
    <header>
        <h1>Test Document Title</h1>
        <p>Author: Test Author</p>
    </header>
    
    <main>
        <section>
            <h2>Introduction</h2>
            <p>This is a test HTML document for Ragwell.</p>
            <p>It contains various HTML elements for testing.</p>
        </section>
        
        <section>
            <h2>Sample Table</h2>
            <table>
                <thead>
                    <tr>
                        <th>Name</th>
                        <th>Age</th>
                        <th>City</th>
                    </tr>
                </thead>
                <tbody>
                    <tr>
                        <td>Alice</td>
                        <td>30</td>
                        <td>New York</td>
                    </tr>
                    <tr>
                        <td>Bob</td>
                        <td>25</td>
                        <td>San Francisco</td>
                    </tr>
                    <tr>
                        <td>Charlie</td>
                        <td>35</td>
                        <td>Chicago</td>
                    </tr>
                </tbody>
            </table>
        </section>
    </main>
    
    <footer>
        <p>End of test document</p>
    </footer>
</body>
</html>"""
    
    output_path.write_text(html_content, encoding='utf-8')


def create_csv_fixture(output_path: Path, delimiter: str = ','):
    """Create a CSV test file."""
    data = [
        ["Name", "Age", "City"],
        ["Alice", "30", "New York"],
        ["Bob", "25", "San Francisco"],
        ["Charlie", "35", "Chicago"]
    ]
    
    with open(output_path, 'w', newline='', encoding='utf-8') as f:
        writer = csv.writer(f, delimiter=delimiter)
        writer.writerows(data)


def create_tsv_fixture(output_path: Path):
    """Create a TSV test file."""
    create_csv_fixture(output_path, delimiter='\t')


def create_txt_fixture(output_path: Path, encoding: str = 'utf-8'):
    """Create a TXT test file."""
    content = """Test Document Title

Author: Test Author

Introduction
============

This is a test text document for Ragwell.
It contains multiple paragraphs and sections.

Section 1
---------

This is the first section of the test document.
It has some sample content for testing purposes.

Section 2
---------

This is the second section.
More test content here.

End of document."""
    
    output_path.write_text(content, encoding=encoding)


def create_all_fixtures():
    """Create all test fixture files."""
    fixtures_dir = Path(__file__).parent / 'fixtures'
    fixtures_dir.mkdir(exist_ok=True)
    
    # Create PDF
    create_pdf_fixture(fixtures_dir / 'test.pdf')
    
    # Create DOCX
    create_docx_fixture(fixtures_dir / 'test.docx')
    
    # Create HTML
    create_html_fixture(fixtures_dir / 'test.html')
    
    # Create CSV
    create_csv_fixture(fixtures_dir / 'test.csv')
    
    # Create TSV
    create_tsv_fixture(fixtures_dir / 'test.tsv')
    
    # Create TXT
    create_txt_fixture(fixtures_dir / 'test.txt')
    
    # Create UTF-16 TXT
    create_txt_fixture(fixtures_dir / 'test_utf16.txt')
    # Re-encode to UTF-16
    content = (fixtures_dir / 'test_utf16.txt').read_text(encoding='utf-8')
    (fixtures_dir / 'test_utf16.txt').write_text(content, encoding='utf-16')
    
    print(f"Created test fixtures in: {fixtures_dir}")
    return fixtures_dir


if __name__ == "__main__":
    create_all_fixtures()
